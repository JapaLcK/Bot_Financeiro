"""
Export por email: POST /export/{uid} gera PDF + XLSX + CSV do período e envia pro
email cadastrado via send_email (com anexos base64). Testa a lógica do handler
com auth mockada e send_email espionado — não toca rede.
"""
from __future__ import annotations

import asyncio
import base64
import io
from datetime import date, datetime

import pytest
from fastapi import HTTPException

import frontend.finance_bot_websocket_custom as app_mod


@pytest.fixture
def _no_auth(monkeypatch):
    import core.services.plan_service as plan_service

    monkeypatch.setattr(app_mod, "_authorize_dashboard_access", lambda req, user_id: None)
    monkeypatch.setattr(app_mod, "_require_pro", lambda user_id, feature: None)
    monkeypatch.setattr(plan_service, "history_earliest_date", lambda _uid, _now=None: None)
    # chamamos export_email direto (sem Request HTTP); desliga o rate-limit do slowapi
    monkeypatch.setattr(app_mod.limiter, "enabled", False)


@pytest.fixture
def spy_email(monkeypatch):
    import core.services.email_service as es
    captured: dict = {}

    def fake_send(to, subject, html_body, text_body=None, from_addr=None, headers=None, attachments=None):
        captured.update(to=to, subject=subject, attachments=attachments)
        return True

    monkeypatch.setattr(es, "send_email", fake_send)
    return captured


def _add_launch(
    uid: int,
    tipo: str,
    valor: float,
    alvo: str,
    categoria: str,
    criado_em: datetime | None = None,
) -> None:
    from db.connection import get_conn
    with get_conn() as conn:
        with conn.cursor() as cur:
            cur.execute(
                "insert into launches (user_id, tipo, valor, alvo, nota, categoria, criado_em) "
                "values (%s, %s, %s, %s, '', %s, coalesce(%s, now()))",
                (uid, tipo, valor, alvo, categoria, criado_em),
            )
        conn.commit()


def test_export_email_envia_3_anexos(_no_auth, spy_email, pro_user_id):
    _add_launch(pro_user_id, "receita", 3000, "Salário", "salário")
    _add_launch(pro_user_id, "despesa", 120.5, "Mercado", "alimentação")

    now = datetime.now()
    res = asyncio.run(app_mod.export_email(None, pro_user_id, now.year, now.month))

    assert res["ok"] is True
    assert "***@" in res["email"]

    atts = spy_email["attachments"]
    assert {a["filename"].rsplit(".", 1)[1] for a in atts} == {"pdf", "xlsx", "csv"}
    by_ext = {a["filename"].rsplit(".", 1)[1]: base64.b64decode(a["content"]) for a in atts}
    assert by_ext["pdf"][:4] == b"%PDF"
    assert by_ext["xlsx"][:2] == b"PK"  # xlsx é um zip


def test_export_email_404_sem_lancamentos(_no_auth, spy_email, pro_user_id):
    now = datetime.now()
    with pytest.raises(HTTPException) as exc:
        asyncio.run(app_mod.export_email(None, pro_user_id, now.year + 1, 1))
    assert exc.value.status_code == 404
    assert spy_email.get("attachments") is None


def test_export_email_aceita_periodo_e_inclui_resumo(
    _no_auth, spy_email, pro_user_id, monkeypatch,
):
    from openpyxl import load_workbook

    balance_calls = 0
    original_fetch_balance = app_mod._fetch_export_balance

    async def counted_fetch_balance(user_id):
        nonlocal balance_calls
        balance_calls += 1
        return await original_fetch_balance(user_id)

    monkeypatch.setattr(app_mod, "_fetch_export_balance", counted_fetch_balance)

    _add_launch(
        pro_user_id, "receita", 3000, "Salário", "salário",
        datetime(2026, 1, 31, 12, 0),
    )
    _add_launch(
        pro_user_id, "despesa", 120.5, "Mercado", "alimentação",
        datetime(2026, 2, 15, 12, 0),
    )
    # Fora do recorte inclusivo: não pode contaminar os totais.
    _add_launch(
        pro_user_id, "despesa", 999, "Fora", "outros",
        datetime(2026, 3, 1, 0, 0),
    )

    res = asyncio.run(app_mod.export_email(
        None,
        pro_user_id,
        start_date=date(2026, 1, 31),
        end_date=date(2026, 2, 28),
    ))

    assert res["ok"] is True
    assert "31/01/2026 a 28/02/2026" in spy_email["subject"]
    atts = {a["filename"].rsplit(".", 1)[1]: a for a in spy_email["attachments"]}
    assert all("20260131_a_20260228" in a["filename"] for a in atts.values())

    csv_text = base64.b64decode(atts["csv"]["content"]).decode("utf-8")
    assert "Salário" in csv_text
    assert "Mercado" in csv_text
    assert "Fora" not in csv_text

    xlsx = base64.b64decode(atts["xlsx"]["content"])
    wb = load_workbook(io.BytesIO(xlsx), data_only=True)
    assert wb.sheetnames[:2] == ["Resumo", "Lançamentos"]
    resumo = {row[0]: row[1] for row in wb["Resumo"].iter_rows(min_row=2, values_only=True)}
    assert resumo["Período"] == "31/01/2026 a 28/02/2026"
    assert resumo["Total de entradas"] == 3000
    assert resumo["Total de saídas"] == 120.5
    assert resumo["Saldo do período"] == 2879.5
    assert resumo["Balanço (receitas - despesas)"] == 2879.5
    assert "Saldo atual" in resumo
    assert resumo["Lançamentos"] == 2
    assert balance_calls == 1


def test_export_email_rejeita_periodo_invertido(_no_auth, spy_email, pro_user_id):
    with pytest.raises(HTTPException) as exc:
        asyncio.run(app_mod.export_email(
            None,
            pro_user_id,
            start_date=date(2026, 3, 1),
            end_date=date(2026, 2, 28),
        ))

    assert exc.value.status_code == 400
    assert "data final" in exc.value.detail
    assert spy_email == {}


def test_export_email_rejeita_data_final_que_estoura_limite(
    _no_auth, spy_email, pro_user_id,
):
    with pytest.raises(HTTPException) as exc:
        asyncio.run(app_mod.export_email(
            None,
            pro_user_id,
            start_date=date(9999, 12, 1),
            end_date=date.max,
        ))

    assert exc.value.status_code == 400
    assert "31/12/9999" in exc.value.detail
    assert spy_email == {}


def test_export_email_limita_inicio_ao_historico_do_plano(
    _no_auth, spy_email, pro_user_id, monkeypatch,
):
    import core.services.plan_service as plan_service

    monkeypatch.setattr(
        plan_service,
        "history_earliest_date",
        lambda _uid, _now=None: date(2026, 2, 1),
    )
    _add_launch(
        pro_user_id, "receita", 3000, "Antes do plano", "salário",
        datetime(2026, 1, 31, 12, 0),
    )
    _add_launch(
        pro_user_id, "despesa", 120.5, "Dentro do plano", "alimentação",
        datetime(2026, 2, 15, 12, 0),
    )

    asyncio.run(app_mod.export_email(
        None,
        pro_user_id,
        start_date=date(2026, 1, 1),
        end_date=date(2026, 2, 28),
    ))

    assert "01 a 28/02/2026" in spy_email["subject"]
    csv_attachment = next(
        a for a in spy_email["attachments"] if a["filename"].endswith(".csv")
    )
    csv_text = base64.b64decode(csv_attachment["content"]).decode("utf-8")
    assert "Antes do plano" not in csv_text
    assert "Dentro do plano" in csv_text


def test_resumo_separa_fluxo_de_caixa_do_balanco_operacional():
    def item(natureza, sign, valor):
        return {
            "natureza": natureza,
            "sign": sign,
            "valor": valor,
            "categoria": "",
        }

    summary = app_mod._export_summary([
        item("receita", "+", 1000),
        item("despesa", "-", 200),
        item("aporte", "-", 300),
        item("aporte", "+", 50),
    ])

    assert summary["entradas"] == 1050
    assert summary["saidas"] == 500
    assert summary["saldo_periodo"] == 550
    assert summary["balanco"] == 800
